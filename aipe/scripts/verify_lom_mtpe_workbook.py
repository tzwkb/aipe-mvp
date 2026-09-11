"""Verify that the LOM MTPE workbook changes only the declared D:E:F cells."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from copy import copy
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from openpyxl import load_workbook


EXPECTED_SOURCE_SHA256 = (
    "f6dd829b6dc2a3b736004d9a4cbf68cd0b4326f7ef3fdef54766a635b47e6489"
)
ALLOWED_ROWS = [
    *range(14, 20),
    *range(22, 28),
    30,
    *range(33, 67),
]
ALLOWED_CELLS = {
    f"{column}{row}" for row in ALLOWED_ROWS for column in ("D", "E", "F")
}
ALLOWED_PACKAGE_CHANGES = {
    "xl/sharedStrings.xml",
    "xl/styles.xml",
    "xl/worksheets/sheet2.xml",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"JSON top level must be an object: {path}")
    return value


def expected_outputs(paths: list[Path]) -> dict[str, Any]:
    outputs: dict[str, Any] = {}
    for path in paths:
        value = read_json(path)
        if value.get("status") != "final":
            raise ValueError(f"translation part is not final: {path}")
        for unit in value.get("units", []):
            row = unit["row"]
            for column, key in (
                ("D", "translation"),
                ("E", "translation_note"),
                ("F", "query"),
            ):
                coordinate = f"{column}{row}"
                if coordinate in outputs:
                    raise ValueError(f"duplicate output coordinate: {coordinate}")
                outputs[coordinate] = unit.get(key, "")
    if set(outputs) != ALLOWED_CELLS:
        raise ValueError("translation output coordinates do not match the 47-unit contract")
    return outputs


def compare_dimensions(source: Any, output: Any, sheet_name: str) -> None:
    if source.max_row != output.max_row or source.max_column != output.max_column:
        raise ValueError(f"worksheet dimensions changed: {sheet_name}")
    if set(map(str, source.merged_cells.ranges)) != set(
        map(str, output.merged_cells.ranges)
    ):
        raise ValueError(f"merged cells changed: {sheet_name}")
    if source.freeze_panes != output.freeze_panes:
        raise ValueError(f"freeze panes changed: {sheet_name}")
    if source.sheet_state != output.sheet_state:
        raise ValueError(f"sheet state changed: {sheet_name}")
    if len(source._images) != len(output._images):
        raise ValueError(f"image count changed: {sheet_name}")
    if len(source.data_validations.dataValidation) != len(
        output.data_validations.dataValidation
    ):
        raise ValueError(f"data validations changed: {sheet_name}")

    if set(source.column_dimensions) != set(output.column_dimensions):
        raise ValueError(f"column dimension keys changed: {sheet_name}")
    for key in source.column_dimensions:
        left = source.column_dimensions[key]
        right = output.column_dimensions[key]
        numeric = ("width",)
        scalar = ("hidden", "outlineLevel", "bestFit")
        for attribute in numeric:
            left_value = getattr(left, attribute)
            right_value = getattr(right, attribute)
            if left_value is None or right_value is None:
                if left_value != right_value:
                    raise ValueError(f"column {key} {attribute} changed: {sheet_name}")
            elif abs(left_value - right_value) > 1e-6:
                raise ValueError(f"column {key} {attribute} changed: {sheet_name}")
        for attribute in scalar:
            if getattr(left, attribute) != getattr(right, attribute):
                raise ValueError(f"column {key} {attribute} changed: {sheet_name}")
        if left._style != right._style:
            raise ValueError(f"column {key} effective style changed: {sheet_name}")

    if set(source.row_dimensions) != set(output.row_dimensions):
        raise ValueError(f"row dimension keys changed: {sheet_name}")
    for key in source.row_dimensions:
        left = source.row_dimensions[key]
        right = output.row_dimensions[key]
        if left.height is None or right.height is None:
            if left.height != right.height:
                raise ValueError(f"row {key} height changed: {sheet_name}")
        elif abs(left.height - right.height) > 1e-6:
            raise ValueError(f"row {key} height changed: {sheet_name}")
        for attribute in ("hidden", "outlineLevel"):
            if getattr(left, attribute) != getattr(right, attribute):
                raise ValueError(f"row {key} {attribute} changed: {sheet_name}")
        if left._style != right._style:
            raise ValueError(f"row {key} effective style changed: {sheet_name}")


def compare_cells(
    source: Any,
    serialization_baseline: Any,
    output: Any,
    sheet_name: str,
    expected: dict[str, Any],
) -> tuple[int, int, list[str]]:
    def hyperlink_signature(value: Any) -> tuple[Any, ...] | None:
        if value is None:
            return None
        return (value.target, value.location, value.tooltip, value.display)

    def comment_signature(value: Any) -> tuple[Any, ...] | None:
        if value is None:
            return None
        return (value.text, value.author)

    def verify_query_style_patch(baseline_cell: Any, output_cell: Any) -> None:
        color = output_cell.font.color
        if color is None or color.type != "rgb" or color.rgb not in {
            "00000000",
            "FF000000",
        }:
            raise ValueError(f"query font is not black: {output_cell.coordinate}")
        if output_cell.alignment.wrap_text is not True:
            raise ValueError(f"query wrap is not enabled: {output_cell.coordinate}")
        if (
            copy(output_cell.fill) != copy(baseline_cell.fill)
            or copy(output_cell.border) != copy(baseline_cell.border)
            or output_cell.number_format != baseline_cell.number_format
            or copy(output_cell.protection) != copy(baseline_cell.protection)
        ):
            raise ValueError(f"query non-font style changed: {output_cell.coordinate}")
        for attribute in (
            "name",
            "sz",
            "b",
            "i",
            "u",
            "strike",
            "vertAlign",
            "charset",
            "family",
            "scheme",
            "outline",
            "shadow",
            "condense",
            "extend",
        ):
            if getattr(output_cell.font, attribute) != getattr(
                baseline_cell.font, attribute
            ):
                raise ValueError(
                    f"query font attribute changed: {output_cell.coordinate}:{attribute}"
                )
        for attribute in (
            "horizontal",
            "vertical",
            "text_rotation",
            "shrink_to_fit",
            "indent",
            "relativeIndent",
            "justifyLastLine",
            "readingOrder",
        ):
            if getattr(output_cell.alignment, attribute) != getattr(
                baseline_cell.alignment, attribute
            ):
                raise ValueError(
                    f"query alignment changed: {output_cell.coordinate}:{attribute}"
                )

    checked = 0
    changed = 0
    style_changes: list[str] = []
    max_row = max(source.max_row, serialization_baseline.max_row, output.max_row)
    max_column = max(
        source.max_column,
        serialization_baseline.max_column,
        output.max_column,
    )
    for row in range(1, max_row + 1):
        for column in range(1, max_column + 1):
            source_cell = source.cell(row, column)
            baseline_cell = serialization_baseline.cell(row, column)
            output_cell = output.cell(row, column)
            coordinate = source_cell.coordinate
            allowed = sheet_name == "试译" and coordinate in ALLOWED_CELLS
            if allowed:
                expected_value = expected[coordinate]
                actual_value = output_cell.value
                if expected_value == "" and actual_value is None:
                    actual_value = ""
                if actual_value != expected_value:
                    raise ValueError(f"unexpected output value at {coordinate}")
                if output_cell.data_type == "f":
                    raise ValueError(f"formula introduced at {coordinate}")
                changed += 1
            else:
                if output_cell.value != source_cell.value:
                    raise ValueError(f"non-output value changed: {sheet_name}!{coordinate}")
                if output_cell.data_type != baseline_cell.data_type:
                    raise ValueError(f"cell type changed: {sheet_name}!{coordinate}")
            if output_cell._style != baseline_cell._style:
                is_readable_query = (
                    sheet_name == "试译"
                    and coordinate.startswith("F")
                    and bool(expected.get(coordinate))
                )
                if not is_readable_query:
                    raise ValueError(f"cell style changed: {sheet_name}!{coordinate}")
                verify_query_style_patch(baseline_cell, output_cell)
                style_changes.append(coordinate)
            if hyperlink_signature(output_cell.hyperlink) != hyperlink_signature(
                baseline_cell.hyperlink
            ):
                raise ValueError(f"hyperlink changed: {sheet_name}!{coordinate}")
            if comment_signature(output_cell.comment) != comment_signature(
                baseline_cell.comment
            ):
                raise ValueError(f"comment changed: {sheet_name}!{coordinate}")
            checked += 1
    return checked, changed, style_changes


def compare_package(baseline_path: Path, output_path: Path) -> list[str]:
    relationship_id = re.compile(rb"R[0-9A-Fa-f]{16}")

    def normalized(name: str, value: bytes) -> bytes:
        if name.endswith((".xml", ".rels")):
            return relationship_id.sub(b"RELATIONSHIP_ID", value)
        return value

    with ZipFile(baseline_path) as baseline_zip, ZipFile(output_path) as output_zip:
        baseline_names = set(baseline_zip.namelist())
        output_names = set(output_zip.namelist())
        if baseline_names != output_names:
            raise ValueError("OOXML package member set changed")
        changed = sorted(
            name
            for name in baseline_names
            if normalized(name, baseline_zip.read(name))
            != normalized(name, output_zip.read(name))
        )
    if set(changed) - ALLOWED_PACKAGE_CHANGES:
        raise ValueError(f"unexpected OOXML package changes: {changed}")
    if "xl/worksheets/sheet2.xml" not in changed:
        raise ValueError("translation worksheet XML did not change")
    return changed


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--serialization-baseline", type=Path, required=True)
    parser.add_argument("--non-dialogue", type=Path, required=True)
    parser.add_argument("--dialogue", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--audit-output", type=Path)
    args = parser.parse_args()

    source = args.source.resolve()
    baseline = args.serialization_baseline.resolve()
    output = args.output.resolve()
    if sha256(source) != EXPECTED_SOURCE_SHA256:
        raise ValueError("source workbook SHA mismatch")
    expected = expected_outputs(
        [args.non_dialogue.resolve(), args.dialogue.resolve()]
    )

    source_workbook = load_workbook(source, data_only=False)
    baseline_workbook = load_workbook(baseline, data_only=False)
    output_workbook = load_workbook(output, data_only=False)
    if not (
        source_workbook.sheetnames
        == baseline_workbook.sheetnames
        == output_workbook.sheetnames
    ):
        raise ValueError("worksheet names/order changed")

    checked_cells = 0
    changed_cells = 0
    style_changes: list[str] = []
    for sheet_name in source_workbook.sheetnames:
        source_sheet = source_workbook[sheet_name]
        baseline_sheet = baseline_workbook[sheet_name]
        output_sheet = output_workbook[sheet_name]
        compare_dimensions(baseline_sheet, output_sheet, sheet_name)
        checked, changed, changed_styles = compare_cells(
            source_sheet,
            baseline_sheet,
            output_sheet,
            sheet_name,
            expected,
        )
        checked_cells += checked
        changed_cells += changed
        style_changes.extend(changed_styles)

    package_changes = compare_package(baseline, output)
    nonempty_queries = sorted(
        coordinate
        for coordinate, value in expected.items()
        if coordinate.startswith("F") and value
    )
    result = {
        "status": "verified",
        "source_sha256": sha256(source),
        "serialization_baseline_sha256": sha256(baseline),
        "output_sha256": sha256(output),
        "worksheets": output_workbook.sheetnames,
        "checked_cells": checked_cells,
        "declared_output_cells": changed_cells,
        "non_output_value_changes": 0,
        "style_changes": len(style_changes),
        "intentional_query_readability_style_changes": sorted(style_changes),
        "merged_cell_changes": 0,
        "row_column_dimension_changes": 0,
        "image_count_changes": 0,
        "unexpected_package_changes": 0,
        "package_changes": package_changes,
        "nonempty_query_cells": nonempty_queries,
    }
    rendered_result = json.dumps(result, ensure_ascii=False, indent=2)
    if args.audit_output:
        audit_output = args.audit_output.resolve()
        audit_output.parent.mkdir(parents=True, exist_ok=True)
        audit_output.write_text(f"{rendered_result}\n", encoding="utf-8")
    print(rendered_result)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
