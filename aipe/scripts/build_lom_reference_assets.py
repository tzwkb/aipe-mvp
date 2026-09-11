"""Build deterministic, local-only LOM reference-corpus evidence assets."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import statistics
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


WORD_RE = re.compile(r"\b[A-Za-z]+(?:[’'][A-Za-z]+)?\b")
SENTENCE_RE = re.compile(r"(?<=[.!?…])\s+")


def _atomic_text(path: Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(value, encoding="utf-8")
    temporary.replace(path)


def _write_json(path: Path, value: Any) -> None:
    _atomic_text(path, json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def _write_jsonl(path: Path, values: list[dict[str, Any]]) -> None:
    _atomic_text(
        path,
        "".join(json.dumps(value, ensure_ascii=False, separators=(",", ":")) + "\n" for value in values),
    )


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _termbase_entries(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    entries: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            try:
                header = [_clean(value) for value in next(rows)]
            except StopIteration:
                continue
            normalized = {name.strip().lower(): index for index, name in enumerate(header) if name}
            source_index = next(
                (normalized[key] for key in ["中文", "原文", "source", "cn"] if key in normalized),
                0,
            )
            target_index = next(
                (normalized[key] for key in ["英文", "译文", "target", "en"] if key in normalized),
                1,
            )
            for row_number, row in enumerate(rows, 2):
                source = _clean(row[source_index] if source_index < len(row) else None)
                target = _clean(row[target_index] if target_index < len(row) else None)
                if not source or not target:
                    continue
                entries.append(
                    {
                        "source": source,
                        "target": target,
                        "category": sheet.title,
                        "coordinate": f"{sheet.title}!{row_number}",
                    }
                )
    finally:
        workbook.close()
    return entries


def _term_pattern(term: str) -> re.Pattern[str]:
    escaped = re.escape(term)
    left = r"(?<![A-Za-z0-9])" if term[:1].isalnum() else ""
    right = r"(?![A-Za-z0-9])" if term[-1:].isalnum() else ""
    return re.compile(left + escaped + right, flags=re.IGNORECASE)


def _line_excerpt(line: str, match_start: int, match_end: int, radius: int = 110) -> str:
    start = max(0, match_start - radius)
    end = min(len(line), match_end + radius)
    prefix = "…" if start else ""
    suffix = "…" if end < len(line) else ""
    return prefix + line[start:end].strip() + suffix


def build_terminology_evidence(
    termbase: Path,
    source: Path,
    chapters: list[dict[str, Any]],
) -> tuple[dict[str, Any], dict[str, list[str]]]:
    entries = _termbase_entries(termbase)
    lines = source.read_text(encoding="utf-8-sig").replace("\r\n", "\n").replace("\r", "\n").splitlines()
    chapter_for_line: list[int] = [0] * (len(lines) + 1)
    for chapter in chapters:
        for line_number in range(int(chapter["line_start"]), int(chapter["line_end"]) + 1):
            chapter_for_line[line_number] = int(chapter["chapter_number"])

    unique: dict[tuple[str, str, str], dict[str, Any]] = {}
    for entry in entries:
        key = (entry["source"], entry["target"], entry["category"])
        current = unique.setdefault(
            key,
            {
                "source": entry["source"],
                "target": entry["target"],
                "category": entry["category"],
                "termbase_coordinates": [],
            },
        )
        current["termbase_coordinates"].append(entry["coordinate"])

    evidence_entries: list[dict[str, Any]] = []
    target_to_sources: dict[str, list[str]] = defaultdict(list)
    for value in unique.values():
        target_to_sources[value["target"].casefold()].append(value["source"])
        pattern = _term_pattern(value["target"])
        occurrences: list[dict[str, Any]] = []
        variants: Counter[str] = Counter()
        for line_number, line in enumerate(lines, 1):
            for match in pattern.finditer(line):
                actual = match.group(0)
                variants[actual] += 1
                occurrences.append(
                    {
                        "chapter": chapter_for_line[line_number] or None,
                        "line": line_number,
                        "matched": actual,
                        "excerpt": _line_excerpt(line, match.start(), match.end()),
                    }
                )
        chapters_seen = sorted({item["chapter"] for item in occurrences if item["chapter"]})
        evidence_entries.append(
            {
                **value,
                "occurrence_count": len(occurrences),
                "chapter_count": len(chapters_seen),
                "first_chapter": chapters_seen[0] if chapters_seen else None,
                "last_chapter": chapters_seen[-1] if chapters_seen else None,
                "surface_variants": dict(variants.most_common()),
                "occurrences": occurrences,
                "verification_status": "corpus_attested" if occurrences else "not_attested_in_lord_txt",
            }
        )

    evidence_entries.sort(key=lambda item: (item["category"], item["source"], item["target"]))
    report = {
        "schema": "aipe.reference-terminology-evidence",
        "version": 1,
        "source": {
            "path": str(source),
            "sha256": _hash(source),
            "lines": len(lines),
        },
        "termbase": {
            "path": str(termbase),
            "sha256": _hash(termbase),
            "row_records": len(entries),
            "unique_source_target_category": len(evidence_entries),
        },
        "coverage": {
            "attested_entries": sum(item["occurrence_count"] > 0 for item in evidence_entries),
            "unattested_entries": sum(item["occurrence_count"] == 0 for item in evidence_entries),
            "total_occurrences": sum(item["occurrence_count"] for item in evidence_entries),
        },
        "entries": evidence_entries,
    }
    return report, target_to_sources


def build_chunk_alias_map(
    chunks: list[dict[str, Any]],
    terminology: dict[str, Any],
) -> list[dict[str, Any]]:
    searchable = [
        (entry, _term_pattern(entry["target"]))
        for entry in terminology["entries"]
        if entry["target"]
    ]
    records: list[dict[str, Any]] = []
    for chunk in chunks:
        matches: list[dict[str, str]] = []
        for entry, pattern in searchable:
            if pattern.search(chunk["text"]):
                matches.append(
                    {
                        "source": entry["source"],
                        "target": entry["target"],
                        "category": entry["category"],
                    }
                )
        aliases = sorted({item["source"] for item in matches})
        records.append(
            {
                "chunk_id": chunk["chunk_id"],
                "chapter_number": chunk["chapter_number"],
                "line_start": chunk["line_start"],
                "line_end": chunk["line_end"],
                "client_terms": matches,
                "source_aliases": aliases,
                "sparse_search_text": " ".join([*aliases, chunk["chapter_title"], chunk["text"]]),
            }
        )
    return records


def _chapter_style(chapter: dict[str, Any]) -> dict[str, Any]:
    paragraphs = [line.strip() for line in chapter["text"].splitlines() if line.strip()]
    prose = [
        line for line in paragraphs
        if not line.startswith("Translator:") and not line.startswith("Editor:")
    ]
    sentences = [
        sentence.strip()
        for paragraph in prose
        for sentence in SENTENCE_RE.split(paragraph)
        if sentence.strip()
    ]
    sentence_words = [len(WORD_RE.findall(sentence)) for sentence in sentences]
    dialogue = [line for line in prose if line.startswith(("“", '"'))]
    first_person = [line for line in prose if re.search(r"\b(I|I’m|I’ve|I’ll|me|my|mine)\b", line)]
    past_markers = sum(len(re.findall(r"\b(?:was|were|had|did|said|asked|looked|walked|felt|saw)\b", line, re.I)) for line in prose)
    present_markers = sum(len(re.findall(r"\b(?:am|is|are|have|do|say|ask|look|walk|feel|see)\b", line, re.I)) for line in prose)
    return {
        "chapter_number": chapter["chapter_number"],
        "title": chapter["title"],
        "line_start": chapter["line_start"],
        "line_end": chapter["line_end"],
        "word_count": chapter["word_count"],
        "paragraph_count": len(prose),
        "dialogue_paragraph_share": round(len(dialogue) / len(prose), 6) if prose else 0,
        "first_person_paragraph_share": round(len(first_person) / len(prose), 6) if prose else 0,
        "sentence_word_median": statistics.median(sentence_words) if sentence_words else 0,
        "past_marker_count": past_markers,
        "present_marker_count": present_markers,
        "punctuation": {
            "curly_open_quote": chapter["text"].count("“"),
            "curly_close_quote": chapter["text"].count("”"),
            "ellipsis": chapter["text"].count("…"),
            "em_dash": chapter["text"].count("—"),
            "exclamation": chapter["text"].count("!"),
            "question": chapter["text"].count("?"),
        },
    }


def build_style_statistics(chapters: list[dict[str, Any]]) -> dict[str, Any]:
    chapter_stats = [_chapter_style(chapter) for chapter in chapters]
    mainline = [item for item in chapter_stats if not str(item["title"]).startswith("In Modern Day")]
    modern = [item for item in chapter_stats if str(item["title"]).startswith("In Modern Day")]

    def aggregate(values: list[dict[str, Any]]) -> dict[str, Any]:
        if not values:
            return {}
        total_paragraphs = sum(item["paragraph_count"] for item in values)
        return {
            "chapters": len(values),
            "words": sum(item["word_count"] for item in values),
            "paragraphs": total_paragraphs,
            "dialogue_paragraph_share_weighted": round(
                sum(item["dialogue_paragraph_share"] * item["paragraph_count"] for item in values)
                / max(total_paragraphs, 1),
                6,
            ),
            "first_person_paragraph_share_weighted": round(
                sum(item["first_person_paragraph_share"] * item["paragraph_count"] for item in values)
                / max(total_paragraphs, 1),
                6,
            ),
            "sentence_word_median_across_chapters": statistics.median(item["sentence_word_median"] for item in values),
            "past_marker_count": sum(item["past_marker_count"] for item in values),
            "present_marker_count": sum(item["present_marker_count"] for item in values),
            "punctuation": {
                key: sum(item["punctuation"][key] for item in values)
                for key in values[0]["punctuation"]
            },
        }

    blocks: list[dict[str, Any]] = []
    for start in range(1, 1429, 100):
        values = [item for item in chapter_stats if start <= item["chapter_number"] <= min(start + 99, 1428)]
        blocks.append({"chapter_start": start, "chapter_end": min(start + 99, 1428), **aggregate(values)})
    return {
        "schema": "aipe.reference-style-statistics",
        "version": 1,
        "overall": aggregate(chapter_stats),
        "mainline": aggregate(mainline),
        "in_modern_day": aggregate(modern),
        "blocks": blocks,
        "chapters": chapter_stats,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--termbase", type=Path, required=True)
    parser.add_argument("--chapters", type=Path, required=True)
    parser.add_argument("--chunks", type=Path, required=True)
    parser.add_argument("--out-dir", type=Path, required=True)
    args = parser.parse_args()

    chapters = _read_jsonl(args.chapters)
    chunks = _read_jsonl(args.chunks)
    terminology, _ = build_terminology_evidence(args.termbase, args.source, chapters)
    alias_map = build_chunk_alias_map(chunks, terminology)
    style = build_style_statistics(chapters)
    args.out_dir.mkdir(parents=True, exist_ok=True)
    _write_json(args.out_dir / "terminology_evidence.json", terminology)
    _write_jsonl(args.out_dir / "chunk_alias_map.jsonl", alias_map)
    _write_json(args.out_dir / "style_statistics.json", style)
    summary = {
        "chapters": len(chapters),
        "chunks": len(chunks),
        "terminology_entries": len(terminology["entries"]),
        **terminology["coverage"],
    }
    _write_json(args.out_dir / "deterministic_asset_summary.json", summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
